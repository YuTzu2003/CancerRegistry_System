import logging
import os
import re
import shutil
from copy import copy
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from modules.blueprint.clean.field_mapping import get_field_map
from modules.services.db import get_conn
from modules.blueprint.clean.cleaner import FORMAT_RULES_MAP

DUPLICATE_COMPARISON_FIELDS = ( "病歷號碼", "身分證統一編號", "癌症發生順序號碼", "最初診斷日期", "原發部位",)

def get_duplicate_comparison_columns(fmt_name, headers, alias_to_target=None):
    fmt_key = f"fmt_{str(fmt_name).replace('fmt_', '')}"
    rules = FORMAT_RULES_MAP.get(fmt_key, {})
    expected = {
        field_name: f"{rules[field_name]['ID']}{field_name}"
        for field_name in DUPLICATE_COMPARISON_FIELDS
        if field_name in rules
    }
    if len(expected) != len(DUPLICATE_COMPARISON_FIELDS):
        raise ValueError("目前資料格式缺少重複資料比對欄位")

    alias_to_target = alias_to_target if alias_to_target is not None else get_field_map("field_name_zh", fmt_name)
    normalized_targets = {re.sub(r"\s+", "", str(target)): field_name for field_name, target in expected.items()}
    resolved = {}
    for header in headers:
        text = str(header).strip()
        target = alias_to_target.get(text, text)
        field_name = normalized_targets.get(re.sub(r"\s+", "", str(target)))
        if field_name and field_name not in resolved:
            resolved[field_name] = header
    missing = [field_name for field_name in DUPLICATE_COMPARISON_FIELDS if field_name not in resolved]
    if missing:
        raise ValueError(f"缺少重複資料比對欄位：{'、'.join(missing)}")
    return resolved

def normalize_duplicate_value(value):
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", "", str(value)).upper()

def get_duplicate_keys(dataframe, columns):
    keys = set()
    for _, row in dataframe.iterrows():
        key = tuple(normalize_duplicate_value(row[columns[field_name]]) for field_name in DUPLICATE_COMPARISON_FIELDS)
        if all(key):
            keys.add(key)
    return keys

def append_reviewed_rows_to_cleaning_result(output_file, reviewed_output_file):
    if not reviewed_output_file or not os.path.exists(reviewed_output_file):
        return

    reviewed_record_marker = "E:資料已清洗過"
    reviewed_record_fill = PatternFill("solid", fgColor="E2F0D9")
    workbook = load_workbook(output_file)
    reviewed_workbook = load_workbook(reviewed_output_file)
    worksheet = workbook.active
    reviewed_worksheet = reviewed_workbook.active
    headers = [cell.value for cell in worksheet[1]]
    reviewed_columns = {str(cell.value): index for index, cell in enumerate(reviewed_worksheet[1], start=1)}
    marker_column = next((index for index, header in enumerate(headers, start=1) if str(header).startswith("錯誤註記說明")),None,)
    if marker_column is None:
        marker_column = len(headers) + 1
        worksheet.cell(row=1, column=marker_column, value="錯誤註記說明")
        headers.append("錯誤註記說明")

    for reviewed_row in reviewed_worksheet.iter_rows(min_row=2):
        excel_row = worksheet.max_row + 1
        for column_index in range(1, len(headers) + 1):
            cell = worksheet.cell(row=excel_row, column=column_index)
            source_column = reviewed_columns.get(str(headers[column_index - 1]))
            if source_column:
                source_cell = reviewed_row[source_column - 1]
                cell.value = source_cell.value
                cell.number_format = source_cell.number_format
                cell.fill = copy(source_cell.fill) if source_cell.fill.fill_type == "solid" else reviewed_record_fill
            else:
                cell.fill = reviewed_record_fill

        marker_cell = worksheet.cell(row=excel_row, column=marker_column)
        marker_text = str(marker_cell.value or "").strip()
        marker_cell.value = f"{marker_text} {reviewed_record_marker}".strip()
    workbook.save(output_file)
    workbook.close()
    reviewed_workbook.close()

def get_review_records_logic(user_id):
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("""SELECT Job.JobID, Job.FileName, Job.TotalCount, Job.CreatedAt, DataFormat.FmtName, DataFormat.Version
                    FROM Job JOIN DataFormat ON Job.FmtID = DataFormat.FmtID WHERE Job.UserID = ? AND Job.CreatedAt >= DATEADD(month, -3, GETDATE())
                    ORDER BY Job.CreatedAt DESC""",(user_id,),)
    records = [
        {
            "job_id": str(row.JobID),
            "file_name": row.FileName,
            "total_count": row.TotalCount,
            "created_at": row.CreatedAt.strftime("%Y/%m/%d %H:%M") if row.CreatedAt else "",
            "format": str(row.FmtName),
            "version": row.Version,
        }
        for row in cursor.fetchall()]
    conn.close()
    return {"ok": True, "records": records}, 200

def cleanup_expired_review_records():
    conn = get_conn()
    cursor = conn.cursor()
    cursor.execute("SELECT JobID, [Path] FROM Job WHERE CreatedAt < DATEADD(month, -3, GETDATE())")
    expired_records = cursor.fetchall()
    if not expired_records:
        conn.close()
        return 0

    job_ids = [row.JobID for row in expired_records]
    placeholders = ", ".join("?" for job_id in job_ids)
    cursor.execute(f"DELETE FROM Job WHERE JobID IN ({placeholders})", job_ids)
    conn.commit()
    conn.close()
    jobs_root = os.path.abspath(os.path.join("tasks", "Jobs"))
    for job_id, project_path in expired_records:
        path = os.path.abspath(str(project_path or ""))
        try:
            if path and os.path.commonpath((jobs_root, path)) == jobs_root and os.path.isdir(path):
                shutil.rmtree(path)
        except OSError:
            logging.exception("Unable to delete expired review record files: %s", path)
    logging.info("Deleted %s expired review records", len(expired_records))
    return len(expired_records)

def get_selected_review_duplicate_keys(user_id, review_job_ids):
    unique_job_ids = list(dict.fromkeys(job_id for job_id in review_job_ids if job_id))
    if not unique_job_ids:
        return set()

    conn = get_conn()
    cursor = conn.cursor()
    placeholders = ", ".join("?" for job_id in unique_job_ids)
    cursor.execute(f"""SELECT Job.JobID, Job.Path, Job.FileName, DataFormat.FmtName FROM Job JOIN DataFormat ON Job.FmtID = DataFormat.FmtID
                    WHERE Job.UserID = ? AND Job.CreatedAt >= DATEADD(month, -3, GETDATE())
                    AND Job.JobID IN ({placeholders})""",[user_id, *unique_job_ids],)
    review_records = cursor.fetchall()
    conn.close()
    if len(review_records) != len(unique_job_ids):
        raise ValueError("選取的資料審核紀錄不存在或已超過保存期限")

    reference_keys = set()
    for record in review_records:
        base_name = os.path.splitext(record.FileName)[0]
        cleaned_file = os.path.join(record.Path, f"fmt{record.FmtName}_{base_name}_Clean.xlsx")
        if not os.path.exists(cleaned_file):
            raise ValueError(f"找不到資料審核紀錄的清洗檔案：{record.FileName}")
        reference_df = pd.read_excel(cleaned_file, dtype=str)
        reference_columns = get_duplicate_comparison_columns(record.FmtName, reference_df.columns)
        reference_keys.update(get_duplicate_keys(reference_df, reference_columns))
    return reference_keys