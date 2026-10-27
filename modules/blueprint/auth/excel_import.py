from openpyxl import Workbook, load_workbook
import xlrd


def load_import_workbook(file_source):
    if not file_source.filename.lower().endswith(".xls"):
        return load_workbook(file_source, read_only=True, data_only=True)

    try:
        source = xlrd.open_workbook(file_contents=file_source.read())
    except xlrd.XLRDError as error:
        raise ValueError("無法讀取 Excel 檔案。") from error
    workbook = Workbook()
    try:
        for index, sheet in enumerate(source.sheets()):
            worksheet = workbook.active if index == 0 else workbook.create_sheet()
            worksheet.title = sheet.name
            for row_index in range(sheet.nrows):
                values = []
                for cell in sheet.row(row_index):
                    value = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
                        value = int(value)
                    values.append(value)
                worksheet.append(values)
        return workbook
    finally:
        source.release_resources()
