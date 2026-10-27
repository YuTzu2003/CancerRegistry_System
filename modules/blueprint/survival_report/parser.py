"""Readers and normalizer for downloaded survival-rate workbooks."""

from __future__ import annotations

import re
from dataclasses import replace
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from .models import (
    AnnualObservation,
    IssueSeverity,
    ParseIssue,
    ParsedWorkbook,
    SourceFormat,
    StageSeries,
    SurvivalEstimate,
)


class UnsupportedSourceFormat(ValueError):
    """Raised when an upload is not one of the supported spreadsheet forms."""


_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
_PERIOD_PATTERN = re.compile(r"(?P<start>\d{4})\s*年至\s*(?P<end>\d{4})\s*年間")
_YEAR_PATTERN = re.compile(r"(?P<year>\d+)\s*年")
_CI_PATTERN = re.compile(
    r"[（(]\s*(-?\d+(?:[.,]\d+)?)\s*[,，~～至]\s*(-?\d+(?:[.,]\d+)?)\s*[）)]"
)


def detect_source_format(path: Path) -> SourceFormat:
    prefix = path.read_bytes()[:512]
    if prefix.startswith(b"PK\x03\x04"):
        return SourceFormat.OOXML
    if prefix.startswith(_OLE_MAGIC):
        return SourceFormat.BIFF
    if prefix.lstrip().startswith(b"<"):
        return SourceFormat.HTML
    raise UnsupportedSourceFormat("不支援的檔案格式；請上傳 .xls 或 .xlsx 存活率報表。")


def _clean_text(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value).replace("\u3000", " ")).strip()


def _read_ooxml_rows(path: Path) -> list[list[object | None]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_biff_rows(path: Path) -> list[list[object | None]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = workbook.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    finally:
        workbook.release_resources()


def _decode_html(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp950", "big5"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _read_html_rows(path: Path) -> list[list[object | None]]:
    soup = BeautifulSoup(_decode_html(path.read_bytes()), "html.parser")
    table = soup.find("table")
    if table is None:
        raise ValueError("未找到可用的存活率資料。")
    rows: list[list[object | None]] = []
    spans: dict[int, tuple[int, str]] = {}
    for tr in table.find_all("tr"):
        row: list[object | None] = []
        column = 0

        def fill_spans() -> None:
            nonlocal column
            while column in spans:
                remaining, text = spans[column]
                row.append(text)
                if remaining <= 1:
                    del spans[column]
                else:
                    spans[column] = (remaining - 1, text)
                column += 1

        for cell in tr.find_all(["th", "td"], recursive=False):
            fill_spans()
            text = cell.get_text(" ", strip=True)
            colspan = max(1, int(cell.get("colspan", 1)))
            rowspan = max(1, int(cell.get("rowspan", 1)))
            for offset in range(colspan):
                row.append(text if offset == 0 else "")
                if rowspan > 1:
                    spans[column] = (rowspan - 1, text if offset == 0 else "")
                column += 1
        fill_spans()
        if row:
            rows.append(row)
    return rows


def _read_rows(path: Path, source_format: SourceFormat) -> list[list[object | None]]:
    if source_format is SourceFormat.OOXML:
        return _read_ooxml_rows(path)
    if source_format is SourceFormat.BIFF:
        return _read_biff_rows(path)
    return _read_html_rows(path)


def _extract_period(*texts: str) -> tuple[int, int]:
    for text in texts:
        match = _PERIOD_PATTERN.search(text)
        if match:
            return int(match.group("start")), int(match.group("end"))
    raise ValueError("無法辨識統計期間（例如 2019年至2023年間）。")


def _strip_report_name(value: str) -> str:
    name = Path(value).stem
    name = _PERIOD_PATTERN.sub("", name)
    name = re.sub(r"^\s*\d{10}-\d+\s*[+＋]?\s*", "", name)
    name = re.sub(r"[-－]?\s*存活率統計.*$", "", name)
    return name.strip(" +-－_()（）")


def _extract_header_hospital(text: str) -> str:
    if "存活率統計" not in text:
        return ""
    return _strip_report_name(text)


def _document_text(rows: Iterable[list[object | None]], limit: int = 10) -> str:
    parts: list[str] = []
    for row_index, row in enumerate(rows):
        if row_index >= limit:
            break
        parts.extend(_clean_text(value) for value in row if _clean_text(value))
    return " ".join(parts)


def _parse_decimal(value: object | None) -> tuple[Decimal | None, str]:
    text = _clean_text(value).replace("％", "").replace("%", "")
    if not text:
        return None, "missing"
    if text in {"-", "—", "–", "－", "N/A", "NA", "不適用"}:
        return None, "not_available" if text not in {"不適用"} else "not_applicable"
    if "," in text and "." not in text and text.count(",") == 1:
        text = text.replace(",", ".")
    try:
        return Decimal(text), "available"
    except InvalidOperation:
        return None, "invalid"


def _parse_int(value: object | None) -> tuple[int | None, str]:
    decimal, status = _parse_decimal(value)
    if decimal is None:
        return None, status
    if decimal != decimal.to_integral_value():
        return None, "invalid"
    return int(decimal), "available"


def _parse_ci(value: object | None) -> tuple[Decimal | None, Decimal | None, str]:
    text = _clean_text(value)
    if not text:
        return None, None, "missing"
    if text in {"-", "—", "–", "－"}:
        return None, None, "not_available"
    match = _CI_PATTERN.search(text)
    if not match:
        return None, None, "invalid"
    low, low_status = _parse_decimal(match.group(1))
    high, high_status = _parse_decimal(match.group(2))
    if low_status != "available" or high_status != "available":
        return None, None, "invalid"
    return low, high, "available"


def _estimate(value: object | None, ci_value: object | None) -> SurvivalEstimate:
    text = _clean_text(value)
    combined_match = re.match(r"^\s*(-?\d+(?:[.,]\d+)?)\s*([（(].*)$", text)
    if combined_match:
        rate, status = _parse_decimal(combined_match.group(1))
        low, high, ci_status = _parse_ci(combined_match.group(2))
    else:
        rate, status = _parse_decimal(value)
        low, high, ci_status = _parse_ci(ci_value)
    if status != "available":
        return SurvivalEstimate.unavailable(status)
    return SurvivalEstimate(
        rate=rate,
        ci_low=low,
        ci_high=high,
        value_status="available" if ci_status in {"available", "missing"} else "invalid_ci",
    )


def _cell(row: list[object | None], index: int) -> object | None:
    return row[index] if index < len(row) else None


def _locate_columns(rows: list[list[object | None]]) -> tuple[int, dict[str, int]]:
    aliases = {
        "cancer": ("癌症別", "癌別"),
        "staging": ("分期系統", "期別類別", "分期"),
        "stage": ("期別",),
        "year": ("追蹤年", "觀察年數", "存活年"),
        "analyzable": ("可分析", "期別人數"),
        "censored": ("設限",),
        "death": ("死亡",),
        "observed": ("觀察存活率",),
        "adjusted": ("年齡調整",),
        "nationwide": ("全國",),
        "medical": ("醫學中心",),
        "non_medical": ("非醫學中心",),
    }
    for row_index, row in enumerate(rows):
        texts = [_clean_text(value) for value in row]
        columns: dict[str, int] = {}
        for key, needles in aliases.items():
            for index, text in enumerate(texts):
                matched = any(needle in text for needle in needles)
                if key == "stage":
                    matched = text in needles
                if matched:
                    columns[key] = index
                    break
        if {"cancer", "stage", "year", "analyzable", "death"}.issubset(columns):
            return row_index, {
                "cancer": 0,
                "staging": 1,
                "stage": 2,
                "year": 3,
                "analyzable": 4,
                "censored": 5,
                "death": 6,
                "observed": 7,
                "adjusted": 8,
                "nationwide": 9,
                "medical": 10,
                "non_medical": 11,
            }
        if {"stage", "year", "analyzable", "death", "observed"}.issubset(columns):
            columns.setdefault("cancer", 0)
            columns.setdefault("staging", 1)
            columns.setdefault("censored", 5)
            columns.setdefault("adjusted", 8)
            columns.setdefault("nationwide", 9)
            columns.setdefault("medical", 10)
            columns.setdefault("non_medical", 11)
            return row_index, columns
    raise ValueError("找不到存活率報表必要欄位。")


def parse_survival_workbook(path: Path, original_filename: str) -> ParsedWorkbook:
    source_format = detect_source_format(path)
    rows = _read_rows(path, source_format)
    document_text = _document_text(rows)
    period_start, period_end = _extract_period(document_text, original_filename)
    target_years = period_end - period_start + 1
    filename_hospital = _strip_report_name(original_filename)
    header_hospital = ""
    for row in rows[:10]:
        for value in row:
            candidate = _extract_header_hospital(_clean_text(value))
            if candidate:
                header_hospital = candidate
                break
        if header_hospital:
            break
    hospital_name = header_hospital or filename_hospital
    issues: list[ParseIssue] = []
    if header_hospital and filename_hospital and header_hospital != filename_hospital:
        issues.append(
            ParseIssue(
                code="hospital_name_conflict",
                message=f"表頭醫院名稱「{header_hospital}」與檔名「{filename_hospital}」不同，已採用表頭。",
                severity=IssueSeverity.WARNING,
                source_file=original_filename,
            )
        )
    if not hospital_name:
        issues.append(
            ParseIssue(
                code="hospital_name_required",
                message="無法辨識醫院名稱，請在預覽頁確認。",
                severity=IssueSeverity.ERROR,
                source_file=original_filename,
            )
        )

    header_row, columns = _locate_columns(rows)
    series_items: list[StageSeries] = []
    cancer_context = ""
    staging_context = ""
    stage_context = ""
    active_key: tuple[str, str, str] | None = None
    active_observations: list[AnnualObservation] = []
    active_estimates: dict[str, SurvivalEstimate] = {}

    def flush_active() -> None:
        nonlocal active_key, active_observations, active_estimates
        if active_key and active_observations:
            series_items.append(
                StageSeries(
                    cancer=active_key[0],
                    staging_system=active_key[1],
                    stage=active_key[2],
                    observations=tuple(active_observations),
                    adjusted=active_estimates.get("adjusted", SurvivalEstimate.unavailable()),
                    nationwide=active_estimates.get("nationwide", SurvivalEstimate.unavailable()),
                    medical_center=active_estimates.get("medical", SurvivalEstimate.unavailable()),
                    non_medical_center=active_estimates.get("non_medical", SurvivalEstimate.unavailable()),
                )
            )
        active_key = None
        active_observations = []
        active_estimates = {}

    for row_index in range(header_row + 1, len(rows)):
        row = rows[row_index]
        year_text = _clean_text(_cell(row, columns["year"]))
        year_match = _YEAR_PATTERN.fullmatch(year_text)
        if not year_match:
            continue
        cancer_value = _clean_text(_cell(row, columns["cancer"]))
        if cancer_value and not cancer_value.startswith(("(", "（")):
            candidate_cancer = re.split(r"\s*個案數\s*[：:]", cancer_value, maxsplit=1)[0].strip()
            if candidate_cancer:
                cancer_context = candidate_cancer
        staging_value = _clean_text(_cell(row, columns["staging"]))
        stage_value = _clean_text(_cell(row, columns["stage"]))
        staging_context = staging_value or staging_context
        stage_context = stage_value or stage_context
        key = (cancer_context, staging_context, stage_context)
        follow_up_year = int(year_match.group("year"))
        if active_key != key or (follow_up_year == 1 and active_observations):
            flush_active()
            active_key = key

        ci_row = rows[row_index + 1] if row_index + 1 < len(rows) else []

        def ci_for(column_key: str) -> object | None:
            current = _cell(row, columns[column_key])
            if _CI_PATTERN.search(_clean_text(current)):
                return current
            return _cell(ci_row, columns[column_key])
        analyzable, analyzable_status = _parse_int(_cell(row, columns["analyzable"]))
        censored, censored_status = _parse_int(_cell(row, columns["censored"]))
        death, death_status = _parse_int(_cell(row, columns["death"]))
        for field_name, status in (
            ("可分析個案數", analyzable_status),
            ("設限個案數", censored_status),
            ("死亡個案數", death_status),
        ):
            if status == "invalid":
                issues.append(
                    ParseIssue(
                        code="invalid_number",
                        message=f"{field_name}無法解析。",
                        severity=IssueSeverity.ERROR,
                        source_file=original_filename,
                        source_row=row_index + 1,
                    )
                )
        observed = _estimate(_cell(row, columns["observed"]), ci_for("observed"))
        active_observations.append(
            AnnualObservation(
                follow_up_year=follow_up_year,
                analyzable_count=analyzable,
                censored_count=censored,
                death_count=death,
                observed=observed,
                source_row=row_index + 1,
            )
        )
        for key_name in ("adjusted", "nationwide", "medical", "non_medical"):
            if key_name not in active_estimates or not active_estimates[key_name].available:
                estimate = _estimate(_cell(row, columns[key_name]), ci_for(key_name))
                if estimate.available or key_name not in active_estimates:
                    active_estimates[key_name] = estimate

    flush_active()
    if not series_items:
        issues.append(
            ParseIssue(
                code="no_survival_rows",
                message="未找到可用的存活率資料列。",
                severity=IssueSeverity.ERROR,
                source_file=original_filename,
            )
        )

    return ParsedWorkbook(
        source_filename=original_filename,
        source_format=source_format,
        hospital_name=hospital_name,
        period_start=period_start,
        period_end=period_end,
        target_years=target_years,
        stage_series=tuple(series_items),
        issues=tuple(issues),
    )
