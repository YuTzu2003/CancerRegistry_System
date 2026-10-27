"""Template-driven editable PowerPoint export."""

from __future__ import annotations

import math
import os
from io import BytesIO
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .models import ReportBatch, ReportDeck, ReportRow, ReportTable, SurvivalEstimate


HEADER_FILL = RGBColor(198, 217, 241)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
HOSPITAL_BLUE = RGBColor(0, 51, 204)
OBSERVATION_LABEL_BLUE = RGBColor(0, 0, 255)
ACCENT_RED = RGBColor(255, 0, 0)
LOCAL_HOSPITAL_BROWN = RGBColor(124, 75, 35)
PAGE_NUMBER_GRAY = RGBColor(191, 191, 191)
FONT_NAME = "微軟正黑體"
FONT_FILE = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts" / "msjh.ttc"
OVERVIEW_MAX_FONT_SIZE = 20
OVERVIEW_MIN_FONT_SIZE = 10

TITLE_GEOMETRY = (457200, 137160, 8229600, 650000)
PAGE_GEOMETRY = (6553200, 6356350, 2133600, 365125)
NOTES_GEOMETRY = (179512, 6165304, 7848872, 430887)
STANDARD_TABLE_GEOMETRY = (35496, 767436, 8712969, 5253849)

_CHINESE_NUMERALS = {
    1: "一",
    2: "二",
    3: "三",
    4: "四",
    5: "五",
    6: "六",
    7: "七",
    8: "八",
    9: "九",
    10: "十",
}


def _remove_slide(presentation: Presentation, index: int) -> None:
    slide_id = presentation.slides._sldIdLst[index]
    presentation.part.drop_rel(slide_id.rId)
    del presentation.slides._sldIdLst[index]


def _set_border(cell, color: str = "000000", width: str = "12700") -> None:
    properties = cell._tc.get_or_add_tcPr()
    for edge_index, edge in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        existing = properties.find(qn(edge))
        if existing is not None:
            properties.remove(existing)
        line = OxmlElement(edge)
        line.set("w", width)
        line.set("cap", "flat")
        line.set("cmpd", "sng")
        line.set("algn", "ctr")
        solid = OxmlElement("a:solidFill")
        rgb = OxmlElement("a:srgbClr")
        rgb.set("val", color)
        solid.append(rgb)
        line.append(solid)
        dash = OxmlElement("a:prstDash")
        dash.set("val", "solid")
        line.append(dash)
        line.append(OxmlElement("a:round"))
        for end_name in ("a:headEnd", "a:tailEnd"):
            end = OxmlElement(end_name)
            end.set("type", "none")
            end.set("w", "med")
            end.set("len", "med")
            line.append(end)
        # DrawingML requires border elements before the cell fill. PowerPoint
        # ignores otherwise-valid lines when they are appended after solidFill.
        properties.insert(edge_index, line)


def _set_cell(
    cell,
    text: str,
    *,
    size: float = 10,
    bold: bool = False,
    color: RGBColor = BLACK,
    fill: RGBColor = WHITE,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = 9525
    cell.margin_right = 9525
    cell.margin_top = 9525
    cell.margin_bottom = 0
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _set_border(cell)


def _format_number(value: int | None) -> str:
    return "-" if value is None else str(value)


def _format_rate(value: SurvivalEstimate) -> str:
    return "-" if value.rate is None else f"{value.rate:.2f}"


def _format_ci(value: SurvivalEstimate) -> str:
    if value.ci_low is None or value.ci_high is None:
        return "-"
    return f"( {value.ci_low:.2f} , {value.ci_high:.2f} )"


def _set_textbox(shape, text: str, size: float, bold: bool = False) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK


def _set_page_number(shape, page_number: int) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = str(page_number)
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run.font.bold = False
    run.font.color.rgb = PAGE_NUMBER_GRAY


def _set_run_typeface(element, font_name: str) -> None:
    properties = element.find(qn("a:rPr"))
    if properties is None:
        properties = OxmlElement("a:rPr")
        element.insert(0, properties)
    for tag in ("a:latin", "a:ea"):
        typeface = properties.find(qn(tag))
        if typeface is None:
            typeface = OxmlElement(tag)
            properties.append(typeface)
        typeface.set("typeface", font_name)


def _normalize_presentation_font(presentation: Presentation) -> None:
    for slide in presentation.slides:
        for tag in (qn("a:r"), qn("a:fld")):
            for element in slide.element.iter(tag):
                text = element.find(qn("a:t"))
                if text is not None and text.text:
                    _set_run_typeface(element, FONT_NAME)


def _set_rich_cell(
    cell,
    parts: tuple[tuple[str, RGBColor], ...],
    *,
    size: float = 12,
    fill: RGBColor = HEADER_FILL,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = 9525
    cell.margin_right = 9525
    cell.margin_top = 9525
    cell.margin_bottom = 0
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    for text, color in parts:
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = color
    _set_border(cell)


def _set_comparison_banner(cell, table_data: ReportTable) -> None:
    target_years = _CHINESE_NUMERALS.get(table_data.target_years, str(table_data.target_years))
    has_local = any(hospital.is_local for hospital in table_data.hospitals)
    scope_parts = (
        (("台大雲林", LOCAL_HOSPITAL_BROWN), ("與全國", BLACK))
        if has_local
        else (("全國", BLACK),)
    )
    _set_rich_cell(
        cell,
        (
            (f"{table_data.period_start}-{table_data.period_end}", ACCENT_RED),
            (" 存活率：", BLACK),
            *scope_parts,
            (target_years, ACCENT_RED),
            ("年存活率比較（百分比）", BLACK),
        ),
    )


def _set_hospital_header(cell, hospital_name: str, size: float) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = HEADER_FILL
    cell.margin_left = 9525
    cell.margin_right = 9525
    cell.margin_top = 9525
    cell.margin_bottom = 0
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    lines = (
        ((hospital_name, BLACK),),
        (("年齡調整", BLACK), ("(觀察)", OBSERVATION_LABEL_BLUE)),
        (("存活率%", BLACK),),
        (("95% CI", BLACK),),
    )
    for line_index, parts in enumerate(lines):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        for text, color in parts:
            run = paragraph.add_run()
            run.text = text
            run.font.name = FONT_NAME
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = color
    _set_border(cell)


def _table_geometry(row_count: int) -> tuple[int, int, int, int, int, int, int]:
    """Return the template geometry and row heights for the table density."""
    if row_count <= 10:
        left, top, width, height = STANDARD_TABLE_GEOMETRY
        return left, top, width, height, 397813, 1140666, NOTES_GEOMETRY[1]
    if row_count <= 14:
        return 179510, 767436, 8784979, 5515658, 324314, 929921, 6427113
    return 323530, 767437, 8496942, 5625771, 276244, 792090, 6427113


def _apply_column_widths(table, table_width: int, hospital_count: int) -> None:
    flexible_count = hospital_count + 3
    if flexible_count == 4 and table_width == STANDARD_TABLE_GEOMETRY[2]:
        widths = (556551, 608351, 608351, 1734929, 1734929, 1734929, 1734929)
    else:
        fixed_total = min(1773253, table_width // 3)
        fixed_widths = (
            fixed_total * 556551 // 1773253,
            fixed_total * 608351 // 1773253,
        )
        third_fixed = fixed_total - sum(fixed_widths)
        flexible_width = (table_width - fixed_total) // flexible_count
        flexible_widths = [flexible_width] * flexible_count
        flexible_widths[-1] += table_width - fixed_total - (flexible_width * flexible_count)
        widths = (*fixed_widths, third_fixed, *flexible_widths)
    for column, width in zip(table.columns, widths):
        column.width = width


def _set_overview_runs(
    cell,
    parts: tuple[tuple[str, RGBColor], ...],
    *,
    alignment: PP_ALIGN = PP_ALIGN.CENTER,
    size: float = OVERVIEW_MAX_FONT_SIZE,
) -> None:
    """Replace body-cell text while preserving the template cell and bullet styling."""
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    paragraph = frame.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = alignment
    for text, color in parts:
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = False
        run.font.color.rgb = color


def _overview_cancer_font_size(cell, column_width: int, row_height: int, text: str) -> int:
    paragraph = cell.text_frame.paragraphs[0]
    properties = paragraph._p.get_or_add_pPr()
    left_indent = int(properties.get("marL", "0"))
    right_indent = int(properties.get("marR", "0"))
    width_points = (
        column_width
        - cell.margin_left
        - cell.margin_right
        - left_indent
        - right_indent
    ) / 12700
    height_points = (row_height - cell.margin_top - cell.margin_bottom) / 12700

    for size in range(OVERVIEW_MAX_FONT_SIZE, OVERVIEW_MIN_FONT_SIZE - 1, -1):
        if FONT_FILE.exists():
            font = ImageFont.truetype(str(FONT_FILE), size * 4)
            line_count = 1
            line_width = 0.0
            for character in text:
                if character == "\n":
                    line_count += 1
                    line_width = 0.0
                    continue
                character_width = font.getlength(character) / 4
                if line_width and line_width + character_width > width_points:
                    line_count += 1
                    line_width = character_width
                else:
                    line_width += character_width
            ascent, descent = font.getmetrics()
            required_height = line_count * ((ascent + descent) / 4)
        else:
            weighted_length = sum(1 if ord(character) > 127 else 0.55 for character in text)
            characters_per_line = max(1, width_points / size)
            line_count = math.ceil(weighted_length / characters_per_line)
            required_height = line_count * size * 1.35
        if required_height <= height_points:
            return size
    return OVERVIEW_MIN_FONT_SIZE


def _overview_cancer_label(table: ReportTable) -> str:
    label = table.cancer.strip()
    staging = table.staging_system.strip()
    if staging.casefold() == "figo" and "figo" not in label.casefold():
        label = f"{label}FIGO"
    if table.rows and all(
        not table.hospital_value(row, table.primary_hospital_key).estimate.available
        for row in table.rows
    ):
        label = f"{label}(個案數太少，無可採用數據)"
    return label


def _overview_groups(
    tables: tuple[ReportTable, ...],
) -> tuple[tuple[int, int, int, str], ...]:
    grouped: dict[tuple[int, int, int], list[str]] = {}
    for table in tables:
        key = (table.period_start, table.period_end, table.target_years)
        label = _overview_cancer_label(table)
        labels = grouped.setdefault(key, [])
        if label and label not in labels:
            labels.append(label)
    return tuple(
        (start, end, years, "、".join(grouped[(start, end, years)]))
        for start, end, years in sorted(grouped)
    )


def _add_data_slide(
    presentation: Presentation,
    table_data: ReportTable,
    logo: tuple[bytes, int, int, int, int] | None,
    page_number: int,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(*TITLE_GEOMETRY)
    _set_textbox(title, table_data.display_title, 32, False)

    if logo:
        blob, left, top, width, height = logo
        slide.shapes.add_picture(BytesIO(blob), left, top, width, height)

    page = slide.shapes.add_textbox(*PAGE_GEOMETRY)
    _set_page_number(page, page_number)

    row_count = 2 + (2 * len(table_data.rows))
    hospital_count = len(table_data.hospitals)
    column_count = 3 + hospital_count + 3
    (
        table_left,
        table_top,
        table_width,
        table_height,
        first_header_height,
        second_header_height,
        notes_top,
    ) = _table_geometry(row_count)
    table_shape = slide.shapes.add_table(
        row_count,
        column_count,
        table_left,
        table_top,
        table_width,
        table_height,
    )
    table = table_shape.table
    _apply_column_widths(table, table_width, hospital_count)
    table.rows[0].height = first_header_height
    table.rows[1].height = second_header_height
    body_row_count = max(1, row_count - 2)
    remaining_height = table_height - first_header_height - second_header_height
    body_height = remaining_height // body_row_count
    for row_index, row in enumerate(list(table.rows)[2:]):
        row.height = body_height + (remaining_height - (body_height * body_row_count) if row_index == body_row_count - 1 else 0)

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(0, column_count - 1))
    _set_cell(table.cell(0, 0), "期別", size=12, bold=True, fill=HEADER_FILL)
    _set_comparison_banner(table.cell(0, 1), table_data)
    headers = (
        "期別人數",
        "死亡人數",
        "全國\n年齡調整存活率%\n95% CI",
        "醫學中心\n年齡調整存活率%\n95% CI",
        "非醫學中心\n年齡調整存活率%\n95% CI",
    )
    header_size = max(8.0, 12 - max(0, column_count - 7) * 0.7)
    _set_cell(table.cell(1, 1), headers[0], size=header_size, bold=True, fill=HEADER_FILL)
    _set_cell(table.cell(1, 2), headers[1], size=header_size, bold=True, fill=HEADER_FILL)
    for hospital_index, hospital in enumerate(table_data.hospitals, start=3):
        hospital_size = min(header_size, 9.5) if len(hospital.display_name) > 12 else header_size
        _set_hospital_header(table.cell(1, hospital_index), hospital.display_name, hospital_size)
    for offset, text in enumerate(headers[2:], start=3 + hospital_count):
        _set_cell(table.cell(1, offset), text, size=header_size, bold=True, fill=HEADER_FILL)

    for index, row_data in enumerate(table_data.rows):
        rate_row = 2 + (index * 2)
        ci_row = rate_row + 1
        for column_index in range(3):
            table.cell(rate_row, column_index).merge(table.cell(ci_row, column_index))
        _set_cell(table.cell(rate_row, 0), row_data.stage, size=12, bold=True)
        _set_cell(table.cell(rate_row, 1), _format_number(row_data.stage_count), size=12)
        _set_cell(table.cell(rate_row, 2), _format_number(row_data.death_count), size=12)
        hospital_values = tuple(
            table_data.hospital_value(row_data, hospital.key)
            for hospital in table_data.hospitals
        )
        values = (
            *((value.estimate, value.value_type) for value in hospital_values),
            (row_data.nationwide, None),
            (row_data.medical_center, None),
            (row_data.non_medical_center, None),
        )
        for column_index, (estimate, value_type) in enumerate(values, start=3):
            hospital_index = column_index - 3
            is_primary = (
                hospital_index < hospital_count
                and table_data.hospitals[hospital_index].key == table_data.primary_hospital_key
            )
            use_observation_blue = is_primary and value_type == "observed"
            _set_cell(
                table.cell(rate_row, column_index),
                _format_rate(estimate),
                size=12,
                bold=is_primary,
                color=HOSPITAL_BLUE if use_observation_blue else BLACK,
            )
            _set_cell(
                table.cell(ci_row, column_index),
                _format_ci(estimate),
                size=12,
                bold=is_primary,
                color=HOSPITAL_BLUE if use_observation_blue else BLACK,
            )

    notes = slide.shapes.add_textbox(
        NOTES_GEOMETRY[0],
        notes_top,
        NOTES_GEOMETRY[2],
        NOTES_GEOMETRY[3],
    )
    frame = notes.text_frame
    frame.clear()
    frame.word_wrap = True
    note_lines = (
        "(1) 藍字：代表年齡層未滿 10 人，無法進行年齡標準化調整，故呈現觀察存活率及 95% CI。",
        "(2) 存活率顯示「-」代表無足夠個案可追蹤至目標年，或所有個案於目標年前死亡，無法計算。",
    )
    for index, text in enumerate(note_lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT_NAME
        run.font.size = Pt(11)
        run.font.color.rgb = BLACK


def _update_overview(presentation: Presentation, tables: tuple[ReportTable, ...]) -> None:
    if len(presentation.slides) < 2:
        return
    slide = list(presentation.slides)[1]
    overview_table = next((shape.table for shape in slide.shapes if shape.has_table), None)
    if overview_table is None:
        return
    groups = _overview_groups(tables)
    body_rows = len(overview_table.rows) - 1
    for offset in range(body_rows):
        row_index = offset + 1
        if offset >= len(groups):
            for column_index in range(3):
                overview_table.cell(row_index, column_index).text_frame.paragraphs[0].clear()
            continue
        start, end, target_years, cancers = groups[offset]
        years_text = _CHINESE_NUMERALS.get(target_years, str(target_years)) + "年"
        _set_overview_runs(
            overview_table.cell(row_index, 0),
            ((str(start), RGBColor(0, 0, 255)), (f"-{end}", BLACK)),
        )
        _set_overview_runs(
            overview_table.cell(row_index, 1),
            ((years_text, ACCENT_RED),),
        )
        cancer_cell = overview_table.cell(row_index, 2)
        cancer_font_size = _overview_cancer_font_size(
            cancer_cell,
            overview_table.columns[2].width,
            overview_table.rows[row_index].height,
            cancers,
        )
        _set_overview_runs(
            cancer_cell,
            ((cancers, BLACK),),
            alignment=PP_ALIGN.LEFT,
            size=cancer_font_size,
        )


def build_presentation(source: ReportDeck | ReportBatch, template_path: Path) -> BytesIO:
    """Build a presentation while retaining the approved template package/theme."""
    selected = source.ordered_tables if isinstance(source, ReportDeck) else source.selected_tables
    if not selected:
        raise ValueError("沒有可匯出的報表。")
    presentation = Presentation(template_path)
    sample_slide = list(presentation.slides)[2] if len(presentation.slides) > 2 else None
    logo = None
    if sample_slide is not None:
        picture = next((shape for shape in sample_slide.shapes if getattr(shape, "image", None)), None)
        if picture is not None:
            logo = (picture.image.blob, picture.left, picture.top, picture.width, picture.height)
    for index in range(len(presentation.slides) - 1, 1, -1):
        _remove_slide(presentation, index)
    _update_overview(presentation, selected)
    for index, table in enumerate(selected, start=3):
        _add_data_slide(presentation, table, logo, index)
    _normalize_presentation_font(presentation)
    output = BytesIO()
    presentation.save(output)
    output.seek(0)
    return output
