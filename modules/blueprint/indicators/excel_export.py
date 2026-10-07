from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="E7EDF2")
INDICATOR_HEADER_FILL = PatternFill("solid", fgColor="FFF2CC")
INDICATOR_VALUE_FILL = PatternFill("solid", fgColor="E8F5E9")
HEADER_FONT = Font(name="Microsoft JhengHei", size=10, bold=True, color="172B4D")
BODY_FONT = Font(name="Microsoft JhengHei", size=10, color="172B4D")
VALUE_FONT = Font(name="Microsoft JhengHei", size=10, bold=True, color="147D3F")
THIN_BORDER = Side(style="thin", color="D9E0E7")
SECTION_BORDER = Side(style="medium", color="7F8C8D")


def create_indicators_excel(dataframe: pd.DataFrame, original_column_count: int) -> BytesIO:
    """Create the downloadable indicator-detail workbook."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dataframe.to_excel(writer, sheet_name="監測指標明細", index=False)
        worksheet = writer.book["監測指標明細"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = True
        worksheet.row_dimensions[1].height = 28

        for cell in worksheet[1]:
            cell.fill = INDICATOR_HEADER_FILL if cell.column > original_column_count else HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell.border = Border(
                left=SECTION_BORDER if cell.column == original_column_count + 1 else THIN_BORDER,
                right=THIN_BORDER,
                top=THIN_BORDER,
                bottom=THIN_BORDER,
            )

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="center")
                if cell.column > original_column_count:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if cell.value == "V":
                        cell.fill = INDICATOR_VALUE_FILL
                        cell.font = VALUE_FONT
                if cell.column == original_column_count + 1:
                    cell.border = Border(left=SECTION_BORDER)

        for column_index, column_name in enumerate(dataframe.columns, start=1):
            if column_index > original_column_count:
                width = max(14, min(24, len(str(column_name)) + 3))
            else:
                sample_values = dataframe.iloc[:200, column_index - 1].dropna().astype(str)
                longest = max([len(str(column_name)), *(len(value) for value in sample_values)], default=10)
                width = max(10, min(28, longest + 2))
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

        worksheet.print_title_rows = "1:1"
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.page_layout_view = False

    output.seek(0)
    return output
